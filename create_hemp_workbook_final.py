#!/usr/bin/env python3
"""
HEMP Dashboard Workbook Generator - FINAL
Updates with Graduated Enrolment Funnel and Employment Outcomes Design
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Load existing workbook
wb = openpyxl.load_workbook('/Users/user/Documents/CHII_Dashboard/HEMP_Dashboard_Updated.xlsx')

# Update Courses sheet with latest changes
ws_courses = wb['Courses']

# Find and update the Courses section
row_start = None
for row in ws_courses.iter_rows(min_row=1, max_row=ws_courses.max_row):
    if row[0].value and "CHANGES MADE" in str(row[0].value):
        row_start = row[0].row
        break

if row_start:
    # Clear old content and add new specifications
    for row in ws_courses.iter_rows(min_row=row_start, max_row=ws_courses.max_row):
        for cell in row:
            cell.value = None

    # Define styles
    SUBHEADER_FILL = PatternFill(start_color="0C447C", end_color="0C447C", fill_type="solid")
    SUBHEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    SECTION_FILL = PatternFill(start_color="479BD6", end_color="479BD6", fill_type="solid")
    SECTION_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    DATA_FONT = Font(name="Calibri", size=10)
    LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="FF6B7280")
    THIN_BORDER = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    row = row_start

    # Add subheader
    cell = ws_courses.cell(row, 1, "ENROLMENT FUNNEL - Graduated Completion Journey")
    cell.font = SUBHEADER_FONT
    cell.fill = SUBHEADER_FILL
    cell.border = THIN_BORDER
    ws_courses.merge_cells(f'A{row}:C{row}')
    ws_courses.row_dimensions[row].height = 22
    row += 1

    # Add chart specifications
    headers = ["Chart Component", "Specification", "Value"]
    for col, header in enumerate(headers, 1):
        cell = ws_courses.cell(row, col, header)
        cell.font = SECTION_FONT
        cell.fill = SECTION_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws_courses.row_dimensions[row].height = 22
    row += 1

    chart_specs = [
        ["Chart Type", "New Structure", "Enrolment Funnel with Cohort Stacking"],
        ["X-Axis", "Stages", "Enrolled, Completed, Graduated"],
        ["Stacked Data", "Cohort Months", "January, May, September"],
        ["", "", ""],
        ["COLOUR SCHEME (Employment Outcomes):", "", ""],
        ["January Colour", "Dark Navy", "#102C5E"],
        ["May Colour", "Medium Blue", "#479BD6"],
        ["September Colour", "Rose/Pink", "#D17A86"],
        ["", "", ""],
        ["CHART DESIGN ELEMENTS:", "", ""],
        ["Height", "Pixels", "270px"],
        ["Margins", "Top, Right, Bottom, Left", "26, 12, 0, -12"],
        ["Bar Category Gap", "Percentage", "40%"],
        ["Bar Size", "Pixels", "46px"],
        ["Bar Radius", "Rounded Corners", "[4, 4, 0, 0]"],
        ["Grid Style", "Dashed Pattern", "strokeDasharray='3 3'"],
        ["Grid Colour", "Light Grey", "rgba(0,33,71,0.06)"],
        ["Grid Direction", "Vertical", "false"],
        ["XAxis Font", "Size and Weight", "11px, fontWeight 600"],
        ["YAxis Font", "Size", "10px"],
        ["Tooltip Cursor", "Colour", "rgba(0,33,71,0.04)"],
        ["Legend", "Type", "Custom div below chart"],
        ["", "", ""],
        ["YEAR FILTER:", "", ""],
        ["Filter Options", "Years", "All Years, 2022, 2023, 2024, 2025, 2026"],
        ["Filter Behaviour", "Updates", "All three bars adjust to selected year"],
        ["Default", "Selection", "All Years (aggregated)"],
        ["", "", ""],
        ["DATA DISTRIBUTION:", "", ""],
        ["January Distribution", "Percentage", "35% of total"],
        ["May Distribution", "Percentage", "40% of total"],
        ["September Distribution", "Percentage", "25% of total"],
        ["", "", ""],
        ["CHANGES FROM PREVIOUS:", "", ""],
        ["Old Metric", "Certification Rate", "Single metric (Certified students)"],
        ["New Metric", "Graduation Rate", "Three-stage funnel (Enrolled→Completed→Graduated)"],
        ["Visualization", "Chart Type", "Previously: Bar chart by year; Now: Stacked by cohort"],
        ["Chart Design", "Reference", "Employment Outcomes by Program (Youth in Work)"],
    ]

    for spec in chart_specs:
        for col, value in enumerate(spec, 1):
            cell = ws_courses.cell(row, col, value)
            cell.font = LABEL_FONT if spec == headers else DATA_FONT
            cell.border = THIN_BORDER
            if value == "" or spec[0] == "":
                cell.font = Font(name="Calibri", size=9, color="#D1D5DB")
            elif any(x in str(value) for x in ["COLOUR", "CHART DESIGN", "YEAR FILTER", "DATA DISTRIBUTION", "CHANGES FROM"]):
                cell.font = Font(name="Calibri", size=10, bold=True, color="#0C447C")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws_courses.row_dimensions[row].height = 20
        row += 1

# Save workbook
wb.save('/Users/user/Documents/CHII_Dashboard/HEMP_Dashboard_Updated.xlsx')

print("✓ HEMP Dashboard Excel Workbook UPDATED!")
print("✓ Location: /Users/user/Documents/CHII_Dashboard/HEMP_Dashboard_Updated.xlsx")
print("✓ Updates Applied:")
print("  ✓ Graduated Enrolment Funnel specifications added")
print("  ✓ Cohort stacking (January, May, September) documented")
print("  ✓ Employment Outcomes chart design specifications")
print("  ✓ Year filter functionality documented")
print("  ✓ Colour scheme from Youth in Work dashboard")
print("  ✓ All chart design elements documented")
