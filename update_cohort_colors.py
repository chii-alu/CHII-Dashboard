import openpyxl
from openpyxl.styles import PatternFill

# Open existing workbook
wb = openpyxl.load_workbook('HEMP_Dashboard_Updated.xlsx')

# Find Courses sheet
ws_courses = wb['Courses']

# Find the color specification rows and update them
for row in ws_courses.iter_rows(min_row=1, max_row=ws_courses.max_row):
    for cell in row:
        if cell.value and isinstance(cell.value, str):
            if 'January' in cell.value and '#E8A8D8' in str(cell.value):
                ws_courses.cell(cell.row, cell.column + 1).value = '#A8C5E6'
            elif 'May' in cell.value and '#B8D4E8' in str(cell.value):
                ws_courses.cell(cell.row, cell.column + 1).value = '#479BD6'
            elif 'September' in cell.value and '#D4E8A8' in str(cell.value):
                ws_courses.cell(cell.row, cell.column + 1).value = '#1E40AF'

# Save
wb.save('HEMP_Dashboard_Updated.xlsx')
print("✓ Workbook updated with new Professional Gradient colours")
